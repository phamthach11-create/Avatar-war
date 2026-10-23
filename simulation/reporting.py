"""Shared JSON, CSV and Excel reporting. No combat rules here."""
import csv
import json
import math
import statistics
from collections import Counter


def summarize(pair, rows):
    wins = Counter(r['winner'] for r in rows)
    n = len(rows); p = wins[pair[0]]/n; z = 1.96
    center = (p+z*z/(2*n))/(1+z*z/n)
    half = z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/(1+z*z/n)
    metrics = {}
    for side, build in enumerate(pair):
        key = 'metrics_a' if side == 0 else 'metrics_b'
        keys = sorted(set().union(*(r[key] for r in rows)))
        metrics[build] = {k:statistics.mean(r[key].get(k,0) for r in rows) for k in keys}
    return dict(pair=pair, wins=dict(wins), a_win_rate=p, a_win_rate_wilson95=[center-half,center+half],
                mean_turns=statistics.mean(r['turns'] for r in rows), median_turns=statistics.median(r['turns'] for r in rows),
                first_wins=sum(r['winner']==r['first_build'] for r in rows),
                by_first={b:dict(games=sum(r['first_build']==b for r in rows),
                                wins=dict(Counter(r['winner'] for r in rows if r['first_build']==b))) for b in pair},
                mean_metrics=metrics)


def export(out, results, summaries, samples, manifest, excel=True):
    from .configuration import write_json
    write_json(out/'summary.json', summaries)
    write_json(out/'manifest.json', manifest)
    with (out/'battles.jsonl').open('w',encoding='utf-8') as f:
        for r in results: f.write(json.dumps(r)+'\n')
    header = ['build_a','build_b','seed','first_build','winner','turns','hp_a','hp_b','mp_a','mp_b']
    with (out/'battles.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f); writer.writerow(header); writer.writerows([[r[k] for k in header] for r in results])
    event_header=['Global turn','Actor build','Event','Detail','Raw damage','Final damage','Effective damage','HP A','HP B','MP A','MP B','AP A','AP B']
    for name, battle in samples:
        with (out/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.writer(f);writer.writerow(event_header);writer.writerows(battle.events)
    if not excel: return
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb=openpyxl.Workbook();sh=wb.active;sh.title='Summary'
    sh.append(['Pair','Build A wins','Build B wins','Draws','A win rate','Mean global turns','First-player wins','95% low','95% high'])
    for s in summaries:
        a,b=s['pair'];sh.append([f'{a} vs {b}',s['wins'].get(a,0),s['wins'].get(b,0),s['wins'].get(0,0),s['a_win_rate'],s['mean_turns'],s['first_wins'],*s['a_win_rate_wilson95']])
        for col in (5,8,9):sh.cell(sh.max_row,col).number_format='0.0%'
    sh=wb.create_sheet('By starting side');sh.append(['Pair','First build','Games','Build A wins','Build B wins','Draws'])
    for s in summaries:
        a,b=s['pair']
        for first,v in s['by_first'].items():sh.append([f'{a} vs {b}',first,v['games'],v['wins'].get(a,0),v['wins'].get(b,0),v['wins'].get(0,0)])
    sh=wb.create_sheet('Average metrics');sh.append(['Pair','Build','Metric','Mean per battle'])
    for s in summaries:
        for b,m in s['mean_metrics'].items():
            for k,v in m.items():sh.append([str(s['pair']),b,k,v])
    sh=wb.create_sheet('Manifest');sh.append(['Setting','Value'])
    for k,v in manifest.items():sh.append([k,json.dumps(v,ensure_ascii=False) if not isinstance(v,str) else v])
    sh=wb.create_sheet('All battles');sh.append(header)
    for r in results:sh.append([r[k] for k in header])
    for name,battle in samples:
        sh=wb.create_sheet(name);sh.append(event_header)
        for row in battle.events:sh.append(row)
    for sh in wb:
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for c in sh[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='234B67')
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=min(65,max(12,max(len(str(c.value or '')) for c in list(col)[:50])+2))
    wb.save(out/'report.xlsx')
