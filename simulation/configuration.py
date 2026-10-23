"""Versioned snapshots and explicit Excel import; never interpret prose as code."""
import copy
import hashlib
import json
import math
from pathlib import Path

ROOT = Path(__file__).resolve().parent
STAT_NAMES = ['HP', 'MP', 'Attack', 'Defense', 'Skill Defense', 'Multi-Hit Rate',
              'Penetration chance', 'Critical Hit Chance', 'Critical Damage Bonus',
              'AP Recovery', 'Stun Threshold', 'Attack stun count', 'Skill Points',
              'Accuracy', 'Evasion', 'Healing Amplification']


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def numeric(value, name, minimum=0):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < minimum:
        raise ValueError(f'{name}: expected finite number >= {minimum}, got {value!r}')
    return value


def assemble_builds(selections, catalog):
    profiles = {}
    for key, selection in selections.items():
        if set(selection) != {'nodes','skill_levels'}:
            raise ValueError('Build selections contain only nodes and skill_levels; numeric effects belong in the node catalog.')
        nodes = selection['nodes']
        if len(nodes) != len(set(nodes)) or len(set(nodes) & {'S1','S2'}) != 1:
            raise ValueError(f'Build {key}: select exactly one specialization and no duplicate nodes')
        profile = copy.deepcopy(selection)
        for node in nodes:
            if node not in catalog: raise ValueError(f'Unsupported passive {node}; add its effect handler before use')
            if not set(catalog[node]['requires']) <= set(nodes): raise ValueError(f'Missing prerequisites for {node}')
            for field,value in catalog[node]['effects'].items():
                if isinstance(value,dict):
                    target=profile.setdefault(field,{})
                    for stat,amount in value.items(): target[stat]=target.get(stat,0)+amount
                elif isinstance(value,(int,float)) and not isinstance(value,bool):
                    profile[field]=profile.get(field,0)+value
                else:
                    profile[field]=copy.deepcopy(value)
        # Decimal percentages in authored JSON should add as decimal percentages.
        for field,value in profile.items():
            if isinstance(value,float): profile[field]=round(value,12)
        profiles[key]=profile
    return profiles


def load_scenario(name='robin_duels'):
    path = Path(name)
    if not path.exists():
        path = ROOT / 'scenarios' / (str(name) + '.json')
    scenario = read_json(path)
    if scenario.get('schema_version') != 1 or scenario.get('engine') not in ('robin_duel','mixed_duel','team_battle'):
        raise ValueError('Unsupported scenario schema/engine; no silent fallback.')
    if scenario['engine'] in ('mixed_duel','team_battle'):
        if scenario['engine']=='mixed_duel' and scenario.get('action_mode') != 'normal_actions':
            raise ValueError('Mixed duels currently support normal voluntary actions plus implemented passives only')
        actors = {}
        rules = read_json(ROOT / scenario['rules'])
        for key, spec in scenario['combatants'].items():
            actor = read_json(ROOT / spec['character'])
            # Experiment-only numeric overrides keep the reviewed source intact.
            for field,value in spec.get('character_overrides',{}).items():
                if field not in STAT_NAMES + ['mark_satk','frozen_dot']:
                    raise ValueError(f'Unsupported character override: {field}')
                actor[field] = numeric(value, field)
            selection = read_json(ROOT / spec['builds'])[spec['build']]
            actor['build_profiles'] = assemble_builds({key: selection}, read_json(ROOT / spec['nodes']))
            actor['rules'] = rules
            actors[key] = actor
        d = dict(id=scenario['id'], combatants=actors, rules=rules,
                 action_mode=scenario['action_mode'],
                 source_sha256={k:v['source_sha256'] for k,v in actors.items()},
                 build_profiles={k:v['build_profiles'][k] for k,v in actors.items()})
        if scenario['engine']=='team_battle':
            d['team_size']=scenario['team_size'];d['main_position']=scenario['main_position']
            d['control_policy']=scenario['control_policy']
            d['victory_condition']=scenario.get('victory_condition','main')
        validate_design(d, scenario)
        return scenario, d
    d = read_json(ROOT / scenario['character'])
    d['build_profiles'] = assemble_builds(read_json(ROOT / scenario['builds']), read_json(ROOT / scenario['nodes']))
    d['rules'] = read_json(ROOT / scenario['rules'])
    validate_design(d, scenario)
    return scenario, d


def validate_design(d, scenario, skip_pair_membership=False):
    if 'combatants' in d:
        if scenario['engine']=='team_battle':
            if d.get('action_mode')!='full_skills' or (d['team_size'],d['main_position']) not in ((9,7),(1,0)) or d['control_policy']!='fresh_row':
                raise ValueError('Team backend supports 9vs9 (position7) or 1vs1 (position0), full skills, fresh-row control policy')
            if d['victory_condition'] not in ('main','elimination'):raise ValueError('Unknown victory condition')
            if scenario['invisible']:raise ValueError('Team Invisible is not implemented for these selected builds')
        elif d.get('action_mode') != 'normal_actions': raise ValueError('Unsupported mixed action mode')
        for pair in scenario['pairs']:
            if any(str(k) not in d['combatants'] for k in pair): raise ValueError('Unknown combatant ID')
        for actor in d['combatants'].values():
            validate_design(actor, scenario, skip_pair_membership=True)
        return
    if d.get('schema_version') != 1:
        raise ValueError('Unsupported character snapshot version')
    for name in STAT_NAMES:
        numeric(d[name], name, 1 if name in ('HP', 'Stun Threshold') else 0)
    numeric(d['reference_skill_level'], 'reference_skill_level', 1)
    numeric(d['mark_satk'], 'mark_satk')
    if d['Attack stun count'] != 0:
        raise ValueError('This backend has not implemented normal-attack stun; add a handler/test before enabling it.')
    if not 0 <= d['Penetration chance'] <= 1:
        raise ValueError('Penetration chance must be between 0 and 1')
    for key, s in d['skills'].items():
        for field in ('mp', 'cd', 'stat'):
            numeric(s[field], f'{key}.{field}')
        if key != 'mark': numeric(s['ap'], f'{key}.ap')
        if key not in ('golden','deadly'): numeric(s['duration'], f'{key}.duration', 1)
        if s['cd'] != int(s['cd']) or (s['duration'] is not None and s['duration'] != int(s['duration'])):
            raise ValueError('Cooldown/duration must be integer global turns')
    rules = d['rules']
    for key in ('power_constant', 'slot_ap_requirement', 'stun_duration', 'stun_count_duration', 'passive_points', 'max_base_skill_level'):
        numeric(rules[key], key, 1)
    for key in ('starting_ap', 'normal_attack_ap', 'defend_ap', 'triggered_golden_ap', 'defend_multiplier'):
        numeric(rules[key], key)
    for key in ('cast_hood_when_ready','mark_amplifies_own_damage'):
        if type(rules[key]) is not bool: raise ValueError(f'{key} must be boolean')
    for key in ('stun_duration','stun_count_duration','passive_points','max_base_skill_level'):
        if rules[key] != int(rules[key]): raise ValueError(f'{key} must be integer')
    if type(scenario['invisible']) is not bool: raise ValueError('invisible must be boolean')
    runs = scenario['runs_per_pair']
    if type(runs) is not int or not 0 < runs <= 1000 or runs % 2:
        raise ValueError('runs_per_pair must be even and between 2 and 1000')
    if type(scenario['turn_limit']) is not int or scenario['turn_limit'] < 1:
        raise ValueError('turn_limit must be a positive integer')
    if not scenario['pairs'] or any(len(p) != 2 or p[0] == p[1] for p in scenario['pairs']):
        raise ValueError('Specify at least one pair of distinct build IDs')
    if len({tuple(sorted(p)) for p in scenario['pairs']}) != len(scenario['pairs']):
        raise ValueError('Duplicate matchup')
    if not skip_pair_membership:
        for pair in scenario['pairs']:
            for b in pair:
                if str(b) not in d['build_profiles']: raise ValueError(f'Unknown build {b}')
    allowed = {'nodes','stat_modifiers','stat_effectiveness','base_bonuses','skills','skill_levels','hood_cd_delta',
               'hood_ap_refund','on_hit_golden_chance','golden_satk_modifier','golden_crit',
               'golden_stun_count','golden_keep_slot_chance','mark_stat_effectiveness','mark_mana_cost','mark_damage','deadly_ignore_condition'}
    supported_stats = set(STAT_NAMES) - {'Skill Points','Attack stun count','Healing Amplification'}
    for key, p in d['build_profiles'].items():
        unknown = set(p)-allowed
        if unknown: raise ValueError(f'Build {key}: unsupported keys {sorted(unknown)}')
        nodes = p['nodes']
        if len(nodes) != len(set(nodes)) or len(nodes) > rules['passive_points']:
            raise ValueError(f'Build {key}: duplicate nodes or passive budget exceeded')
        team=scenario['engine']=='team_battle'
        supported_skills={'golden','mark','frozen','deadly'} if team else {'golden','hood','mark'}
        if ((not p['skills'] or 'golden' not in p['skills']) and scenario.get('action_mode') != 'normal_actions' and not team) or set(p['skills'])-supported_skills:
            raise ValueError(f'Build {key}: unsupported skill kit')
        if team:
            unsupported={'on_hit_golden_chance','golden_crit','golden_stun_count','golden_keep_slot_chance'}
            if any(p.get(k,0) for k in unsupported):raise ValueError('Team backend supports Robin build3 / Bluetooth Striker; additional passive handlers require review')
            if 'weapon' not in d or d['weapon']['target'] not in ('Front','Any') or d['weapon']['area'] not in ('Single','Row Adjacent'):
                raise ValueError('Unsupported team weapon targeting')
        if len(set(p['skills'])) != len(p['skills']) or {'hood','mark'} <= set(p['skills']):
            raise ValueError('A build must select only one specialization skill')
        if set(p['skill_levels']) != set(p['skills']): raise ValueError('Allocate levels to exactly the unlocked skills')
        if sum(p['skill_levels'].values()) > d['Skill Points']: raise ValueError('Skill point budget exceeded')
        for level in p['skill_levels'].values():
            if type(level) is not int or not 1 <= level <= rules['max_base_skill_level']:
                raise ValueError('Skill level outside allocation limit')
        for stat_group in ('stat_modifiers','stat_effectiveness','base_bonuses'):
            for name, value in p.get(stat_group,{}).items():
                if name not in supported_stats: raise ValueError(f'Unsupported build stat {name}')
                numeric(value, name, -1e9)
        for name in supported_stats:
            final = (d[name]+p.get('base_bonuses',{}).get(name,0))*(1+p.get('stat_modifiers',{}).get(name,0))*(1+p.get('stat_effectiveness',{}).get(name,0))
            numeric(final, f'Build{key}.{name}', 1 if name in ('HP','Stun Threshold') else 0)
            if name in ('Critical Hit Chance','Penetration chance') and final > 1:
                raise ValueError(f'Build{key}.{name} exceeds 100%')
        for name in ('on_hit_golden_chance','golden_keep_slot_chance'):
            if not 0 <= numeric(p.get(name,0), name) <= 1: raise ValueError(f'{name} must be 0..1')
        for name in ('hood_ap_refund','golden_stun_count','mark_stat_effectiveness','mark_mana_cost'):
            if name in p: numeric(p[name], name)
        if p.get('golden_stun_count',0) not in (0,1):
            raise ValueError('Current rule permits at most one Stun Count per target per global turn')
        numeric(p.get('golden_satk_modifier',0),'golden_satk_modifier',-1)
        if 'hood' in p['skills'] and d['skills']['hood']['cd']+p.get('hood_cd_delta',0) < 1: raise ValueError('Hood cooldown must remain >=1')


def import_excel(path, template=None):
    """Update numbers only after checking reviewed mechanics and passive prose."""
    import openpyxl
    d = copy.deepcopy(template or read_json(ROOT/'data/characters/robin_hood.json'))
    w = openpyxl.load_workbook(path, data_only=True)
    contract = read_json(ROOT/'data/robin_workbook_contract.json')
    changed = []
    for name, cells in contract.items():
        if name not in w.sheetnames: raise ValueError(f'Missing sheet: {name}')
        for address, reviewed in cells.items():
            if w[name][address].value != reviewed: changed.append(f'{name}!{address}')
    if changed:
        raise ValueError('Mechanic/skill/node text changed; review handler and contract before import: '+', '.join(changed[:15]))
    s = w['Ranger class']
    if s['B2'].value != d['name']: raise ValueError('Character changed; choose an appropriate engine/profile')
    rows = {str(s.cell(r,1).value).strip():r for r in range(9,25)}
    for name in STAT_NAMES:
        if name not in rows: raise ValueError(f'Missing stat: {name}')
        d[name] = numeric(s.cell(rows[name],2).value, name)
    for key, row in [('golden',29),('hood',30),('mark',31)]:
        d['skills'][key] = dict(ap=s.cell(row,2).value, mp=s.cell(row,3).value,
                              cd=s.cell(row,4).value, duration=s.cell(row,8).value, stat=s.cell(row,11).value)
        for field in ('mp','cd','stat'): numeric(d['skills'][key][field],f'{key}.{field}')
    d['mark_satk'] = numeric(s['K32'].value,'mark_satk')
    d['source_sha256'] = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    d['nodes_text'] = {c.coordinate:c.value for row in s.iter_rows(min_row=36,max_row=41) for c in row if c.value is not None}
    return d
