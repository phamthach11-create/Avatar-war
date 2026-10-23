import argparse,csv,json,zipfile
from pathlib import Path
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
parser=argparse.ArgumentParser()
parser.add_argument('--input', type=Path, required=True)
args=parser.parse_args()
p=args.input
summary=json.loads((p/'summary.json').read_text())
w=openpyxl.Workbook();s=w.active;s.title='Results'
s.append(['Starting build','Battles','Build 1 wins','Build 4 wins','Unresolved','Build 4 win rate','Mean global turns','Median global turns'])
for r in summary:s.append([r['first'],r['n'],r['build1_wins'],r['build4_wins'],r['unresolved'],r['build4_wins']/r['n'],r['mean_turns'],r['median_turns']])
for row in s.iter_rows(min_row=2):row[5].number_format='0.0%'
s=w.create_sheet('Mean metrics');s.append(['Starting build','Build','Metric','Mean per battle'])
for r in summary:
 for i,m in enumerate(r['mean_metrics']):
  for k,v in m.items():s.append([r['first'],1 if i==0 else 4,k,v])
s=w.create_sheet('Assumptions');s.append(['Category','Description'])
for k,v in json.loads((Path(__file__).parent/'assumptions.json').read_text()).items():
 for value in v if isinstance(v,list) else [v]:s.append([k,json.dumps(value,ensure_ascii=False) if isinstance(value,dict) else str(value)])
for file,title in [('sample_events','Sample events'),('sample_turns','Sample turns')]:
 s=w.create_sheet(title)
 with (p/(file+'.csv')).open() as f:
  for row in csv.reader(f):
   vals=[]
   for value in row:
    try:vals.append(float(value) if '.' in value else int(value))
    except ValueError:vals.append(value)
   s.append(vals)
s=w.create_sheet('All battles');s.append(['Seed','Starting build','Winner','Global turns','Build 1 survivors','Build 4 survivors','Build 1 main HP','Build 4 main HP'])
for r in json.loads((p/'results.json').read_text()):s.append([r['seed'],r['first'],r['winner'],r['turns'],*r['alive'],*r['main_hp']])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='234E70')
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(75,max(15,max(len(str(c.value or '')) for c in col[:50])+2))
 for row in s.iter_rows(min_row=2):
  for c in row:
   if isinstance(c.value,float) and c.number_format=='General':c.number_format='#,##0.00'
 if s.title=='Assumptions':
  s.column_dimensions['B'].width=110
  for row in s.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True);s.row_dimensions[row[0].row].height=50
out=p/'Bluetooth_9v9_latest.xlsx';w.save(out)
with zipfile.ZipFile(out) as z:assert z.testzip() is None
print('Saved validated report',out)
