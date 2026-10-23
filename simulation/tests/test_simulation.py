import copy
import importlib.util
import tempfile
import unittest
from pathlib import Path

from simulation.cli import execute
from simulation.configuration import ROOT, assemble_builds, import_excel, load_scenario, read_json, validate_design
from simulation.engines.robin_duel import Battle, Fighter
from simulation.reporting import export


class CombatTests(unittest.TestCase):
    def setUp(self):
        self.scenario,self.d=load_scenario()

    def battle(self,pair=(1,2),**kwargs):
        b=Battle(self.d,pair,0,12,**kwargs);b.t=1;return b

    def test_cooldown_starts_on_first_own_turn(self):
        b=self.battle();b.start(0);b.end()
        self.assertEqual(b.f[0].cd['golden'],0)
        self.assertEqual(b.f[1].cd['golden'],1)
        b.t=2;b.start(1);b.end();self.assertEqual(b.f[1].cd['golden'],0)

    def test_trigger_pays_mana_but_no_ap_or_hit(self):
        b=self.battle();u=b.f[0];u.cd['golden']=0
        b.golden(0,triggered=True)
        self.assertEqual((u.ap,u.mp,u.cd['golden']),(0,24350,1))
        self.assertEqual(u.metrics['hit_attempts'],0)

    def test_hood_visibility_override_preserves_other_effects(self):
        b=self.battle(invisible=True);u=b.f[0];u.cd['hood']=0;u.ap=100;b.hood(0)
        self.assertEqual((u.ap,u.mp,u.hood_until,u.cd['hood']),(75,22150,6,8))
        self.assertAlmostEqual(b.chc(0),.69)
        self.assertFalse(b.targetable(1));b.invisible=False;self.assertTrue(b.targetable(1))
        b.invisible=True;b.t=6;self.assertFalse(b.targetable(1))
        b.t=7;self.assertTrue(b.targetable(1));self.assertAlmostEqual(b.chc(0),.6)

    def test_stun_cap_duration_and_passive_exception(self):
        b=self.battle((2,3))
        for t in (1,3,5,7):
            b.t=t;b.add_stun(0);stacks=b.f[1].stacks;b.add_stun(0)
            self.assertEqual(b.f[1].stacks,stacks)
        self.assertEqual((b.f[1].stun_until,b.f[1].stacks),(9,0))
        b.t=8;b.f[1].cd['mark']=0
        self.assertFalse(b.start(1));self.assertEqual(b.f[1].metrics['mark_casts'],1)
        self.assertEqual(b.f[1].ap,0)
        b.t=10;self.assertTrue(b.start(1))

    def test_stun_stack_gap_expires(self):
        b=self.battle((2,3));b.add_stun(0);b.t=3;b.end()
        self.assertEqual(b.f[1].stacks,0)

    def test_preserved_slot_can_spend_into_debt(self):
        b=self.battle();u=b.f[0];u.ap=50
        b.normal(0);self.assertEqual(u.ap,-50)
        b.t=3;self.assertFalse(b.start(0));self.assertEqual(u.ap,50)

    def test_one_initial_slot_not_one_per_100_ap(self):
        b=self.battle();b.f[0].ap=1000
        b.run(limit=1)
        self.assertEqual(b.f[0].metrics['normal_attacks'],1)

    def test_mark_order_is_explicit_and_configurable(self):
        b=self.battle((3,1));b.f[0].cd['mark']=0;b.mark(0)
        self.assertAlmostEqual(b.f[0].metrics['mark_damage'],5760*1000/1880*1.24)
        self.d['rules']['mark_amplifies_own_damage']=False
        b=self.battle((3,1));b.f[0].cd['mark']=0;b.mark(0)
        self.assertAlmostEqual(b.f[0].metrics['mark_damage'],5760*1000/1880)

    def test_nodes_determine_stats_and_skill_levels(self):
        f=Fighter(1,self.d);self.assertAlmostEqual(f.mhr,2.1);self.assertAlmostEqual(f.acc,1.2)
        self.d['build_profiles']['1']['skill_levels']['golden']=4
        self.assertEqual(Fighter(1,self.d).golden_satk,5760)

    def test_unapplied_status_is_not_active_before_first_turn(self):
        b=Battle(self.d,(3,1),0,12,invisible=True)
        self.assertTrue(b.targetable(0));self.assertEqual(b.amp(0),1)
        self.assertEqual(b.chc(0),.6)
        self.assertAlmostEqual(b.mitigation_factor(0),1000/1870)

    def test_invalid_node_budget_and_unimplemented_skill_fail(self):
        catalog=read_json(ROOT/'data/robin_nodes.json')
        with self.assertRaises(ValueError):assemble_builds({'1':{'nodes':['S1','S1.1.4'],'skill_levels':{}}},catalog)
        with self.assertRaises(ValueError):assemble_builds({'1':{'nodes':['S2','S2.2.1'],'skill_levels':{}}},catalog)
        self.d['build_profiles']['1']['skill_levels']['golden']=6
        with self.assertRaises(ValueError):validate_design(self.d,self.scenario)

    def test_1000_run_limit_and_even_initiative(self):
        for runs in (0,999,1002):
            self.scenario['runs_per_pair']=runs
            with self.assertRaises(ValueError):validate_design(self.d,self.scenario)

    def test_full_regression_matches_previous_3000_battles(self):
        rows,summaries,_=execute(self.scenario,self.d)
        baseline=read_json(ROOT/'baselines/robin_no_invisible.json')
        self.assertEqual(len(rows),3000)
        for current,old in zip(summaries,baseline):
            self.assertEqual({str(k):v for k,v in current['wins'].items()},old['wins'])
            self.assertEqual(current['mean_turns'],old['mean_turns'])
        for r in rows:
            for side in ('a','b'):
                m=r['metrics_'+side];opponent='b' if side=='a' else 'a'
                self.assertAlmostEqual(m.get('effective_damage',0),self.d['HP']-r['hp_'+opponent])
                self.assertEqual(m.get('mark_blocked_invisible',0),0)

    def test_export_and_replay_smoke(self):
        import openpyxl
        self.scenario['runs_per_pair']=2;self.scenario['pairs']=[[1,3]]
        self.scenario['invisible']=True
        rows,summaries,samples=execute(self.scenario,self.d)
        with tempfile.TemporaryDirectory() as temp:
            out=Path(temp);export(out,rows,summaries,samples,{'scenario':self.scenario})
            wb=openpyxl.load_workbook(out/'report.xlsx',read_only=True)
            self.assertEqual(wb['All battles'].max_row,3);wb.close()
            self.assertEqual(len(list(out.glob('*first*.csv'))),2)

    def test_excel_numeric_import_and_semantic_change_guard(self):
        import openpyxl
        source=ROOT.parent/'data/excel/Thiết kế game.xlsx'
        d=import_excel(source);self.assertEqual(d['HP'],self.d['HP'])
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'input.xlsx'
            wb=openpyxl.load_workbook(source,data_only=True)
            wb['Ranger class']['B9']=40000;wb.save(path)
            self.assertEqual(import_excel(path)['HP'],40000)
            wb['Ranger class']['H36']='When casting, trigger Golden Arrow';wb.save(path)
            with self.assertRaisesRegex(ValueError,'text changed'):import_excel(path)

    def test_historical_bluetooth_snapshot_replays(self):
        folder=ROOT/'legacy/bluetooth_9v9'
        spec=importlib.util.spec_from_file_location('bluetooth_legacy',folder/'simulate.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        old=read_json(folder/'sample_result.json');d=read_json(folder/'input_stats.json')
        result=module.Battle(d,old['seed'],0 if old['first']==1 else 1).run()
        self.assertEqual(result,old)


if __name__=='__main__':unittest.main()
